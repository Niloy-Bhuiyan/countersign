"""Open a supplier document, or decide it cannot be opened.

Intake never raises on a bad file. A corrupt, encrypted or image-only document is
a normal event in an accounts-payable inbox; failing the batch over one would be
worse than the document itself. Each is given a quarantine status and routed to a
person, and the reason is recorded so the routing report can count it.
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

READABLE = "readable"
UNREADABLE = "quarantined_unreadable"
ENCRYPTED = "quarantined_encrypted"
IMAGE_ONLY = "quarantined_image_only"

#: A spreadsheet that inflates past this when unzipped is treated as unreadable,
#: so a 4 MB upload cannot expand into gigabytes in memory.
MAX_UNZIPPED = 64 * 1024 * 1024


@dataclass
class Intake:
    path: Path
    sha256: str
    media_type: str
    status: str
    #: Text layer for a PDF, one string per page.
    pages: list[str] = field(default_factory=list)
    #: Cell values for a spreadsheet or CSV, one list per row, all as strings.
    rows: list[list[str]] = field(default_factory=list)

    @property
    def readable(self) -> bool:
        return self.status == READABLE


def _cells(values) -> list[str]:
    return ["" if value is None else str(value).strip() for value in values]


def read(path: Path) -> Intake:
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    suffix = path.suffix.lower()

    if suffix == ".csv":
        # Accounting systems still export Windows-1252; UTF-8 first, then that.
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
        try:
            rows = [_cells(row) for row in csv.reader(io.StringIO(text, newline=""))]
        except csv.Error:
            return Intake(path, digest, "text/csv", UNREADABLE)
        return Intake(path, digest, "text/csv", READABLE, rows=rows)

    if suffix == ".xlsx":
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        # Parsers of untrusted files fail in many ways; each one means "a person
        # opens this", never a crashed request.
        try:
            with zipfile.ZipFile(path) as archive:
                if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED:
                    return Intake(path, digest, media, UNREADABLE)
            sheet = load_workbook(path, read_only=True, data_only=True).active
            rows = [_cells(row) for row in sheet.iter_rows(values_only=True)]
        except Exception:  # noqa: BLE001
            return Intake(path, digest, media, UNREADABLE)
        return Intake(path, digest, media, READABLE, rows=rows)

    # PDF. An encrypted file announces itself in the trailer before any parser
    # runs, which lets it be told apart from one that is merely broken.
    if b"/Encrypt" in data:
        return Intake(path, digest, "application/pdf", ENCRYPTED)
    try:
        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
    except Exception:  # noqa: BLE001 - any parser failure on an untrusted file
        return Intake(path, digest, "application/pdf", UNREADABLE)

    if not any(page.strip() for page in pages):
        return Intake(path, digest, "application/pdf", IMAGE_ONLY)
    return Intake(path, digest, "application/pdf", READABLE, pages=pages)
