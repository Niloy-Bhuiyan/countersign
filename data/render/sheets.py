"""Render an invoice as a spreadsheet or a CSV.

Not every supplier sends a PDF. Smaller vendors email a spreadsheet, and the
result is not a clean rectangular table: a header block sits above the lines, the
totals sit below them in the last two columns, and there are blank rows in
between. Anything that assumes ``read_excel`` gives it a dataframe of line items
gets a dataframe of header text instead.

Format is assigned per invoice, deterministically from its id, so the mix is
reproducible.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from data.invoices import Invoice
from data.records import Vendor

#: Roughly 78% PDF, 15% spreadsheet, 7% CSV.
_PDF, _XLSX, _CSV = "pdf", "xlsx", "csv"


def format_for_invoice(invoice: Invoice) -> str:
    bucket = int(invoice.id.split("-")[1]) % 100
    if bucket < 7:
        return _CSV
    if bucket < 22:
        return _XLSX
    return _PDF


def _header_rows(invoice: Invoice, vendor: Vendor) -> list[list[str]]:
    return [
        [vendor.legal_name],
        [f"BIN/TIN: {vendor.tax_id}"],
        [],
        ["Invoice Number", invoice.invoice_number],
        ["Invoice Date", invoice.issued_at.isoformat()],
        ["Payment Due", invoice.due_at.isoformat()],
        ["Purchase Order", invoice.po_id or "-"],
        ["Currency", invoice.currency],
        [],
    ]


COLUMNS = ["Line", "SKU", "Description", "UoM", "Quantity", "Unit Price", "Tax %", "Amount"]


def _line_rows(invoice: Invoice) -> list[list]:
    return [
        [
            line.line_no,
            line.sku,
            line.description,
            line.uom,
            str(line.quantity),
            str(line.unit_price),
            f"{line.tax_rate:g}",
            str(line.line_total),
        ]
        for line in invoice.lines
    ]


def _total_rows(invoice: Invoice) -> list[list]:
    pad = [""] * (len(COLUMNS) - 2)
    return [
        [],
        pad + ["Subtotal", str(invoice.subtotal)],
        pad + ["Tax", str(invoice.tax_total)],
        pad + ["Total Payable", str(invoice.total)],
        [],
        ["Synthetic document generated for the Countersign project. Not a real invoice."],
    ]


def render_csv(invoice: Invoice, vendor: Vendor, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerows(_header_rows(invoice, vendor))
        writer.writerow(COLUMNS)
        writer.writerows(_line_rows(invoice))
        writer.writerows(_total_rows(invoice))
    return path


def render_xlsx(invoice: Invoice, vendor: Vendor, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Invoice"

    for row in _header_rows(invoice, vendor):
        sheet.append(row)
    sheet.append(COLUMNS)
    header_row = sheet.max_row
    for row in _line_rows(invoice):
        sheet.append(row)
    for row in _total_rows(invoice):
        sheet.append(row)

    sheet.cell(row=1, column=1).font = Font(bold=True, size=14)
    for column in range(1, len(COLUMNS) + 1):
        sheet.cell(row=header_row, column=column).font = Font(bold=True)
        sheet.cell(row=header_row, column=column).alignment = Alignment(horizontal="center")
    for column, width in zip("ABCDEFGH", (6, 16, 40, 8, 12, 14, 8, 16), strict=False):
        sheet.column_dimensions[column].width = width

    book.save(path)
    return path
